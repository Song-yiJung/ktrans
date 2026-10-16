#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ktrans 메모 정리 — build_ttl.py 다음 단계

build_ttl.py가 만든 ktrans-data.ttl의 메모(crm:P3_has_note)를
노드당 하나로 합치고, 문장마다 꼬리표를 붙인다.

꼬리표 네 가지
    (설명)  자료에 적힌 사실. UeDeKo 서지·비고, 게재 형태, 역어
    (해석)  누군가의 판단. 판단한 사람의 이름을 문장 안에 쓴다
    (편찬)  데이터를 만든 사람(정송이)의 판단. 레코드 통합, 정규화 근거
    (근거)  문단 번호(para_n)나 문헌

사람이 TTL을 열었을 때 속성 하나(crm:P3_has_note)만 알면 되고,
LLM은 꼬리표로 문장의 성격을 구별해 읽는다. 메모는 한국어(@ko)와 영어(@en)를 함께 둔다.
비평 진술에는 영어 요지(rdfs:label @en)를 더한다 — 영어 대응표는 source/en/ 에 있다.

중역 번역에 붙은 충실도 진술에는 비교 기준을 밝히는 (편찬) 문장을 더한다.
생략이 역자의 선택인지 중간 저본에서 온 것인지 데이터로는 알 수 없기 때문이다.

같은 뜻은 늘 같은 말로 쓴다. 쿼리가 문장을 CONTAINS로 찾기 때문이다.
    "게재 형태는 신문 연재이다" / "잡지 연재" / "잡지 게재" / "잡지 부록" / "단행본"
    "원제의 Leiden을 '…'(으)로 옮겼다"
    "번역유형이 서로 다르다"

사용법
    python add_notes.py ktrans-data.ttl \
        --meta uedeko_werther_318_meta.json --en source/en \
        --wonmun UeDeKo_Werther_원문층.xlsx \
        -o ktrans-data.ttl
"""
import re
import csv
import json
import argparse
from pathlib import Path
from collections import Counter

# ---------------------------------------------------------------- 문구
PERIODICALS = {"每日申報": "신문", "時事評論": "잡지", "新民": "잡지", "文藝月刊": "잡지"}
VENUE_EN = {"每日申報": "Maeil Sinbo (每日申報)", "時事評論": "Sisa Pyeongnon (時事評論)",
            "新民": "Sinmin (新民)", "文藝月刊": "Munye Wolgan (文藝月刊)"}
FORM_EN = {"신문 연재": "newspaper serial", "잡지 연재": "magazine serial", "신문 게재": "single newspaper issue",
           "잡지 게재": "single magazine issue", "잡지 부록": "magazine supplement", "단행본": "book"}
TYPE_EN = {"완역": "complete", "편역": "partial", "개작": "adaptation", "대역본": "bilingual",
           "대역": "bilingual", "중역": "relay", "확인불가": "unverifiable"}
TAG_EN = {"(설명)": "(description)", "(해석)": "(interpretation)", "(편찬)": "(compilation)", "(근거)": "(evidence)"}

BASIS = {
    "KAC": ("국립중앙도서관 저자명전거(KAC)로 정규형을 확인했다.",
            "The normalised name was checked against the National Library of Korea authority file (KAC)."),
    "한글표기": ("UeDeKo에 적힌 한글 표기를 그대로 정규형으로 삼았다.",
             "The Hangul form given in UeDeKo was taken as the normalised name."),
    "괄호병기": ("UeDeKo에 괄호로 병기된 한글 이름을 정규형으로 삼았다.",
             "The Hangul name given in brackets in UeDeKo was taken as the normalised name."),
    "학습쌍": ("같은 데이터 안의 한자·한글 병기 사례에서 한글 이름을 가져왔다.",
            "The Hangul name was taken from a Hanja–Hangul pair elsewhere in the same data."),
    "음추정": ("KAC에 등재되지 않아 한자의 음으로 한글 이름을 추정했다.",
            "Not in KAC; the Hangul name was inferred from the reading of the Hanja."),
    "비평텍스트": ("필명·호를 번역비평 텍스트에 근거해 실명으로 풀었다.",
              "The pen name was resolved to the real name on the basis of the critical essay."),
    "단체명": ("단체 명의이므로 개인으로 정규화하지 않았다.",
            "A corporate name, so it was not normalised to a person."),
}

CRITICS = {
    "kt:P-조향": ("(설명) UeDeKo 『젊은 베르터의 고뇌』 번역비평의 필자이다. 이 데이터의 비평 진술 대부분이 그의 판단이다.",
                "(description) Author of the UeDeKo critical essay on Werther. Most critical statements in this data are this author's judgements."),
    "kt:P-김동소": ("(설명) 조향의 번역비평에 인용된 연구자이다.",
                 "(description) A scholar cited in Jo Hyang's critical essay."),
    "kt:P-김미지": ("(설명) 조향의 번역비평에 인용된 연구자이다. 김미지(2018)는 백화 역의 저본을 확인했다.",
                 "(description) A scholar cited in Jo Hyang's critical essay. Kim Mi-ji (2018) identified the source of Baek-hwa's translation."),
    "kt:P-이유영": ("(설명) 조향의 번역비평에 인용된 연구자이다.",
                 "(description) A scholar cited in Jo Hyang's critical essay."),
    "kt:P-조우호": ("(설명) 조향의 번역비평에 인용된 연구자이다.",
                 "(description) A scholar cited in Jo Hyang's critical essay."),
}

GOETHE = ("(설명) 『젊은 베르터의 고뇌』(1774)의 저자이다. UeDeKo 원저자명 칸에 나오는 표기 26종을 skos:altLabel로 묶었다.\n"
          "(편찬) 원저자명 값 가운데 '베르테르'(1964년 『인기스타아 서한문』)는 작중인물이므로 이 인물에 묶지 않았다.",
          "(description) Author of Die Leiden des jungen Werthers (1774). The 26 forms of his name in UeDeKo's author column are kept as skos:altLabel.\n"
          "(compilation) One author value, '베르테르' (Werther, in a 1964 letter anthology), names the character, so it is not linked to this person.")

# F28 저본 보강 메모 — build_ttl.py의 SOURCE_EVENTS 메모를 꼬리표 문장으로
SOURCE_NOTES = {
    "김동소는 김영보가 읽은 일본어역을 다카야마 초규 역으로 추정 (para 12)": (
        "(해석) 김동소는 김영보가 읽은 일본어역을 다카야마 초규의 번역(1891)으로 추정한다.\n(근거) para_12",
        "(interpretation) Kim Dong-so infers that the Japanese version Kim Yeong-bo read was Takayama Chogyū's (1891).\n(evidence) para_12"),
    "김미지 2018, 342·346 — 궈모뤄 중문판을 저본으로 확인 (para 24)": (
        "(해석) 김미지(2018, 342·346쪽)는 궈모뤄의 중국어역(1922)이 저본임을 확인했다.\n(근거) para_24",
        "(interpretation) Kim Mi-ji (2018, pp. 342, 346) established Guo Moruo's Chinese translation (1922) as the source.\n(evidence) para_24"),
    "역자 서문: 일본역 3~4종과 영역 1종을 대조. 판본 미특정 (para 41)": (
        "(설명) 역자 서문에서 일본어역 3~4종과 영역 1종을 대조했다고 밝혔다. 어느 판본인지는 특정되지 않는다.\n(근거) para_41",
        "(description) In his preface the translator says he compared three or four Japanese versions and one English version. The editions are not identified.\n(evidence) para_41"),
    "저본은 일본어역으로 추정. 판본 미특정 (para 59)": (
        "(해석) 조향은 박용철 역의 저본을 일본어역으로 추정한다. 어느 판본인지는 특정되지 않는다.\n(근거) para_59",
        "(interpretation) Jo Hyang infers a Japanese version as the source of Park Yong-cheol's translation. The edition is not identified.\n(evidence) para_59"),
    "역자가 Jubiläums Ausgabe 16권을 저본으로 명시 (각주)": (
        "(설명) 역자가 Jubiläums Ausgabe 16권을 저본으로 밝혔다.\n(근거) 조향 번역비평의 각주",
        "(description) The translator names Jubiläums-Ausgabe vol. 16 as his source.\n(evidence) a footnote in Jo Hyang's essay"),
    "최초의 독일어 직접 번역으로 추정 (para 88)": (
        "(해석) 조향은 이 번역을 독일어 원문에서 옮긴 첫 번역으로 추정한다.\n(근거) para_88",
        "(interpretation) Jo Hyang infers that this is the first translation made from the German original.\n(evidence) para_88"),
    "독일어로부터의 첫 완역 (para 89)": (
        "(해석) 조향은 이 번역을 독일어 원문에서 옮긴 첫 완역으로 본다.\n(근거) para_89",
        "(interpretation) Jo Hyang regards this as the first complete translation from the German original.\n(evidence) para_89"),
}

E_JA = ("(설명) 원작의 약 4/5 분량을 옮긴 일본어역이다. 『山形日報』에 연재되었다.\n"
        "(해석) 김동소는 이 번역을 김영보 역(1923)의 저본으로 추정한다.\n(근거) para_12",
        "(description) A Japanese translation covering about four fifths of the original, serialised in Yamagata Nippō (山形日報).\n"
        "(interpretation) Kim Dong-so infers that this was the source of Kim Yeong-bo's translation (1923).\n(evidence) para_12")

# 번역이 아닌 표현형의 영어 이름
EXPR_EN = {"kt:E_de": "German original (1774, Weygand)",
           "kt:E_zh": "Guo Moruo's Chinese translation 少年維特之煩惱 (1922, Taidong Shuju)",
           "kt:E_ja_takayama": "Takayama Chogyū's Japanese translation (1891, serialised in Yamagata Nippō)"}

# 중역 번역에 붙은 충실도 진술 — 비교 기준을 밝힌다 (새 속성 없이 메모로)
RELAY_FIDELITY = ("(편찬) 비교 기준은 독일어 원문이다. 이 번역은 중역이므로 생략이 저본에서 온 것인지는 이 데이터로 판단할 수 없다.",
                  "(compilation) The comparison is with the German original. This is a relay translation, so the data cannot tell whether an omission comes from the intermediate source.")
RELAY_EXTRA = {
    "T001": (" 저본으로 추정되는 다카야마 역도 원작의 약 4/5 분량이다(para_12).",
             " The presumed source, Takayama's Japanese version, itself covers about four fifths of the original (para_12)."),
}


# ---------------------------------------------------------------- 도구
def esc(s):
    return str(s).replace("\\", "\\\\").replace('"', '\\"').replace("\r", "")


def note_line(text):
    """여러 줄 메모 — Turtle 긴 문자열. text = (ko, en)"""
    ko, en = text
    return f'    crm:P3_has_note """{esc(ko)}"""@ko ,\n                    """{esc(en)}"""@en'


def lit_one(s):
    return '"' + esc(s).replace("\n", "\\n") + '"'


def ro(word):
    """역어 뒤 조사 (으)로"""
    hanja = {"悲歌": "로", "悲惱": "로", "悲歎": "으로"}
    if word in hanja:
        return hanja[word]
    ch = word[-1]
    if "가" <= ch <= "힣":
        jong = (ord(ch) - 0xAC00) % 28
        return "로" if jong in (0, 8) else "으로"
    return "(으)로"


def ranges(nums):
    nums = sorted(nums)
    out, s, p = [], nums[0], nums[0]
    for n in nums[1:]:
        if n == p + 1:
            p = n
            continue
        out.append(f"{s}" if s == p else f"{s}~{p}")
        s = p = n
    out.append(f"{s}" if s == p else f"{s}~{p}")
    return "·".join(out)


def venue_of(series):
    for k in PERIODICALS:
        if series.startswith(k):
            return k
    return None


def unlit(s):
    """'"..."@ko' 리터럴 -> 문자열"""
    m = re.match(r'\s*"(.*)"(@\w+)?\s*$', s)
    v = m.group(1) if m else s
    return v.replace('\\"', '"').replace("\\n", "\n").replace("\\\\", "\\")


# ---------------------------------------------------------------- 노드별 메모
def en_type(v):
    return "; ".join(TYPE_EN.get(x.strip(), x.strip()) for x in str(v).split(";"))


def to_en(ko):
    """꼬리표만 영어로 바꾼다 — 사람이 따로 옮기지 않은 문장의 비상용"""
    for k, v in TAG_EN.items():
        ko = ko.replace(k, v)
    return ko


def translation_note(tid, old, meta, EN):
    ko, en = [], []
    rec_nums = [int(tid[1:])]
    merge = next((o for o in old if o.startswith("UeDeKo 레코드")), None)
    if merge:
        m = re.match(r"UeDeKo 레코드 ([\d,]+) 통합 \((.*?)\) — (.*)$", merge)
        rec_nums = [int(x) for x in m.group(1).split(",")]
        reason, basis = m.group(2), m.group(3)
    recs = [meta[str(n)] for n in rec_nums if str(n) in meta]
    first = recs[0]
    series = (first.get("총서명") or "").strip()
    pub = (first.get("출판사") or "").strip()
    notes_ue = [r.get("비고", "").strip() for r in recs]
    venue = venue_of(series)

    if venue:
        kind = PERIODICALS[venue]
        form = f"{kind} 연재" if len(recs) > 1 else f"{kind} 게재"
        k = f"(설명) 게재 형태는 {form}이다. 게재처는 『{venue}』" + (f"({pub})이다." if pub else "이다.")
        e = f"(description) Publication form: {FORM_EN[form]}. Venue: {VENUE_EN[venue]}" + (f", {pub}." if pub else ".")
        dates = [r.get("발행연도", "") for r in recs if r.get("발행연도")]
        if len(dates) > 1:
            dates = sorted(dates, key=lambda d: [int(x) for x in re.findall(r"\d+", d)])
            k += f" 게재 시기는 {dates[0]}~{dates[-1]}이다."
            e += f" Published {dates[0]}–{dates[-1]}."
        elif dates:
            k += f" 게재 시기는 {dates[0]}이다."
            e += f" Published {dates[0]}."
    elif any("부록" in n for n in notes_ue):
        form = "잡지 부록"
        k = "(설명) 게재 형태는 잡지 부록이다." + (f" 출판사는 {pub}이다." if pub else "")
        e = "(description) Publication form: magazine supplement." + (f" Publisher: {pub}." if pub else "")
    else:
        form = "단행본"
        k = "(설명) 게재 형태는 단행본이다." + (f" 출판사는 {pub}이다." if pub else "")
        e = "(description) Publication form: book." + (f" Publisher: {pub}." if pub else "")
        if series:
            k += f" 총서는 『{series}』이다."
            e += f" Series: {series}."
        book = (first.get("번역서명") or "").strip()
        if book and book != (first.get("개별작품제목") or "").strip():
            k += f" 수록서는 『{book}』이다."
            e += f" Contained in: {book}."
    ko.append(k); en.append(e)

    uniq = []
    for n in notes_ue:
        if n and n not in uniq:
            uniq.append(n)
    if uniq:
        ko.append("(설명) UeDeKo 비고: " + " / ".join(uniq))
        en.append("(description) UeDeKo remark: " + " / ".join(EN["remark"].get(n, n) for n in uniq))

    if merge:
        ko.append(f"(편찬) UeDeKo 레코드 {ranges(rec_nums)}번({len(rec_nums)}건)을 한 작품으로 통합했다({reason}). {basis}")
        en.append(f"(compilation) UeDeKo records {ranges(rec_nums).replace('~', '–').replace('·', ', ')} ({len(rec_nums)} records) "
                  f"were merged into one work ({EN['merge'].get(reason, reason)}). {EN['merge'].get(basis, basis)}")
    if any(o.startswith("저본유형 ≠ 작품유형") for o in old):
        r = next((r for r in recs if r.get("저본 번역유형") != r.get("작품 번역유형")), first)
        a, b = r.get("저본 번역유형"), r.get("작품 번역유형")
        ko.append(f"(편찬) UeDeKo의 저본 번역유형과 작품 번역유형이 서로 다르다(저본 '{a}', 작품 '{b}').")
        en.append(f"(compilation) UeDeKo's source translation type and work translation type differ (source '{en_type(a)}', work '{en_type(b)}').")
    return ("\n".join(ko), "\n".join(en)), form


def nomen_note(old):
    w = old[0].replace("Leiden 역어: ", "").strip() if old else "미상"
    if w == "미상":
        return ("(설명) 제목에서 원제의 Leiden에 해당하는 역어를 찾을 수 없다.",
                "(description) No word corresponding to Leiden in the original title can be identified.")
    return (f"(설명) 원제의 Leiden을 '{w}'{ro(w)} 옮겼다.",
            f"(description) The title renders Leiden as '{w}'.")


def statement_note(old, para, relay_extra=None):
    quote = next((o.replace("근거표현: ", "", 1) for o in old if o.startswith("근거표현: ")), "")
    ko, en = [], []
    for o in old:
        if o.startswith("값(통제어휘 밖)"):
            v = o.split(":", 1)[1].strip()
            ko.append(f"(편찬) 통제어휘에 없는 값: {v}")
            en.append(f"(compilation) Value outside the controlled vocabulary: {v}")
    if relay_extra is not None:
        ko.append(RELAY_FIDELITY[0] + relay_extra[0])
        en.append(RELAY_FIDELITY[1] + relay_extra[1])
    ko.append(f"(근거) {para}" + (f" “{quote}”" if quote else ""))
    en.append(f"(evidence) {para}" + (f", quoted in Korean: “{quote}”" if quote else ""))
    return "\n".join(ko), "\n".join(en)


def person_note(subj, old):
    if subj == "kt:P-goethe":
        return GOETHE
    if subj in CRITICS:
        return CRITICS[subj]
    if old:
        b = old[0].replace("정규형 판정근거: ", "").strip()
        k, e = BASIS.get(b, (f"정규형 판정 근거: {b}", f"Basis for the normalised name: {b}"))
        return ("(편찬) " + k, "(compilation) " + e)
    return None


# ---------------------------------------------------------------- 본체
def statements(text):
    """최상위 Turtle 문장 단위로 자른다. 리터럴은 한 줄이라는 전제."""
    blocks, cur = [], []
    for line in text.split("\n"):
        if not cur:
            if line and not line[0].isspace() and not line.startswith("#") and not line.startswith("@"):
                cur = [line]
                if line.rstrip().endswith(" ."):
                    blocks.append(("stmt", cur)); cur = []
            else:
                blocks.append(("raw", [line]))
        else:
            cur.append(line)
            if line.rstrip().endswith(" ."):
                blocks.append(("stmt", cur)); cur = []
    if cur:
        blocks.append(("stmt", cur))
    return blocks


def rewrite(stmt, ctx, stats):
    meta, EN, relay = ctx["meta"], ctx["en"], ctx["relay"]
    subj = stmt[0].split()[0]
    head_types = stmt[0]
    body = "\n".join(stmt)
    old, keep = [], []
    for ln in stmt:
        m = re.match(r"\s*crm:P3_has_note\s+(\".*\"(?:@\w+)?)\s*[;.]\s*$", ln)
        if m and ln.startswith("    "):
            old.append(unlit(m.group(1)))
        else:
            keep.append(ln)

    note = None
    if " a ktrans:Translation" in head_types:
        note, form = translation_note(subj.split(":")[1], old, meta, EN)
        stats["게재 형태: " + form] += 1
    elif " a lrmoo:F12_Nomen" in head_types and subj != "kt:nomen-W01-de":
        note = nomen_note(old)
    elif " a crm:E13_Attribute_Assignment" in head_types:
        m = re.search(r"crm:P16_used_specific_object kt:SP(\d+)", body)
        tgt = re.search(r"crm:P140_assigned_attribute_to kt:(T\d+)", body)
        extra = None
        if "crm:P177_assigned_property_of_type kt:dim-fidelity" in body and tgt and tgt.group(1) in relay:
            extra = RELAY_EXTRA.get(tgt.group(1), ("", ""))
            stats["중역 충실도 진술에 비교 기준 메모"] += 1
        note = statement_note(old, f"para_{int(m.group(1))}" if m else "para_?", extra)
        sid = subj.split(":")[1]
        if sid in EN["label"]:
            i = next(k for k, ln in enumerate(keep) if ln.lstrip().startswith("rdfs:label"))
            keep.insert(i + 1, f'    rdfs:label "{esc(EN["label"][sid])}"@en ;')
            stats["영어 요지"] += 1
    elif " a crm:E21_Person" in head_types:
        note = person_note(subj, old)
    elif subj == "kt:E_ja_takayama":
        note = E_JA
    elif old and subj.startswith("kt:ev-"):
        pairs = [SOURCE_NOTES.get(o, ("(설명) " + o, "(description) " + o)) for o in old]
        note = ("\n".join(p[0] for p in pairs), "\n".join(p[1] for p in pairs))
    elif old:
        note = ("\n".join(old), to_en("\n".join(old)))
        stats["처리 안 된 메모: " + subj] += 1

    if subj in EXPR_EN:
        i = next(k for k, ln in enumerate(keep) if ln.lstrip().startswith("rdfs:label"))
        keep.insert(i + 1, f'    rdfs:label "{esc(EXPR_EN[subj])}"@en ;')
    if note is None and not old:
        return keep
    while keep and not keep[-1].strip():
        keep.pop()
    keep[-1] = re.sub(r"\s*[;.]\s*$", "", keep[-1])
    if note is None:
        keep[-1] += " ."
        return keep
    keep[-1] += " ;"
    stats["메모 노드"] += 1
    return keep + [note_line(note) + " ."]


def read_tsv(path, key, val):
    with open(path, encoding="utf-8") as f:
        return {r[key]: r[val] for r in csv.DictReader(f, delimiter="\t")}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("infile")
    ap.add_argument("--meta", required=True)
    ap.add_argument("--en", required=True, help="영어 대응표 폴더 (source/en)")
    ap.add_argument("--wonmun", help="원문층 xlsx — 원문 문단을 전문으로 다시 넣는다")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()

    meta = {r["번호"]: r for r in json.load(open(a.meta, encoding="utf-8"))}
    en_dir = Path(a.en)
    EN = {"label": read_tsv(en_dir / "statement_labels_en.tsv", "id", "en"),
          "remark": read_tsv(en_dir / "uedeko_remarks_en.tsv", "ko", "en"),
          "merge": read_tsv(en_dir / "merge_basis_en.tsv", "ko", "en")}
    text = Path(a.infile).read_text(encoding="utf-8")
    relay = set(re.findall(r"^kt:(T\d+) ktrans:viaType kt:via-relay \.", text, re.M))
    relay |= {t for t, b in re.findall(r"^kt:(T\d+) a ktrans:Translation ;(.*?) \.\n", text, re.M | re.S)
              if "ktrans:viaType kt:via-relay" in b}
    ctx = {"meta": meta, "en": EN, "relay": relay}
    stats = Counter()

    if "### 메모 규칙" not in text:
        rule = ("### 메모 규칙 · Note convention ##################################\n"
                "# 모든 노드의 설명은 crm:P3_has_note 하나에 담는다. 한국어(@ko)와 영어(@en)를 함께 둔다.\n"
                "# Every node carries its context in crm:P3_has_note, in Korean (@ko) and English (@en).\n"
                "# (설명) description — 자료에 적힌 사실 · what the source says\n"
                "# (해석) interpretation — 누군가의 판단, 이름을 함께 쓴다 · someone's judgement, named\n"
                "# (편찬) compilation — 데이터를 만든 사람의 판단 · the compiler's decision\n"
                "# (근거) evidence — 문단 번호나 문헌 · paragraph or reference\n"
                "# 출처 Source: UeDeKo http://www.uedeko.or.kr/wiki2/index.php?title=젊은_베르터의_고뇌_(Die_Leiden_des_jungen_Werthers)\n")
        text = text.replace("\n\n### 작품과 원천 표현형", "\n\n" + rule + "\n### 작품과 원천 표현형", 1)

    out = []
    for kind, lines in statements(text):
        out.extend(rewrite(lines, ctx, stats) if kind == "stmt" else lines)
    text = "\n".join(out)

    if a.wonmun:
        from openpyxl import load_workbook
        ws = load_workbook(a.wonmun, read_only=True)["원문층"]
        full = {int(r[0]): str(r[5] or "") for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        def sub(m):
            pid = int(m.group(1))
            new = lit_one(full.get(pid, ""))
            if new != m.group(2):
                stats["원문 문단 복원"] += 1
            return f'crm:P1_is_identified_by "para_{m.group(1)}" ;\n    crm:P190_has_symbolic_content {new}'
        text = re.sub(r'crm:P1_is_identified_by "para_(\d+)" ;\n    crm:P190_has_symbolic_content (".*")',
                      sub, text)

    Path(a.out).write_text(text, encoding="utf-8")
    for k, v in sorted(stats.items()):
        print(f"  {v:>4}  {k}")


if __name__ == "__main__":
    main()
