#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
원문층 -> 구조층(비평 진술 태깅) 시트 생성

원문층의 서술 문단에서 비평 진술을 뽑기 위한 작업 시트를 만든다.
자동으로 채울 수 있는 칸(para_id, 대상번역본, 진술주체 후보, 근거표현 후보)은
미리 넣어 두고, 판단이 필요한 칸(진술유형·차원·값·극성·요지)은 비워 둔다.

사용법:
    python build_gujocheung.py out/UeDeKo_Werther_원문층.xlsx -o out
    python build_gujocheung.py out/UeDeKo_Werther_원문층.xlsx -o out --all

기본 범위:
    일제강점기 전체 + 1945년 이후의 2.1 개관 절
    (--all 을 주면 전체 문단을 대상으로 한다)

출력:
    UeDeKo_Werther_구조층.xlsx
      - 진술 시트 : 태깅 작업 시트 (드롭다운 포함)
      - 분류표    : 10항목 차원·값 참조
      - 원문참조  : 범위에 든 문단 전문
"""

import re
import argparse
from pathlib import Path
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

PLUM = "6D2E46"
FILL_AUTO = "F2F2F2"      # 자동 채움 칸
FILL_INPUT = "FFFDF0"     # 직접 입력 칸
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FT = "Arial"

# 대상번역본 정규화 — 연도 보완, 시대 구분 제거
TARGET_FIX = {
    "김용호": "김용호(1953)",
    "임홍배": "임홍배(2012)",
    "장만영": "장만영(1951)",
}
NOT_A_TARGET = re.compile(r"^(19\d0|20\d0|\d{4}년|1945년부터|2000년 이후|19\d0-19\d0년대|19\d0년대)")

# 분류표 10항목
DIMENSIONS = [
    ("1", "sentence_ending", "종결어미체", "관찰", "-다 / -네·-데 / -오·-소 / -야·-어 / -습니다", "없음"),
    ("2", "addressee_term", "수신자 호칭", "관찰", "너 / 자네 / 형 / 君 / 호명회피", "없음"),
    ("3", "lexical_register", "어휘 색채", "관찰", "한문투 / 중립 / 현대적", "없음"),
    ("4", "proper_noun", "고유명 처리", "관찰", "원음 / 자국화 / 대체·생략", "없음"),
    ("5", "naturalness", "자연스러움·가독성", "평가", "자연스럽다 / 어색하다 / 건조하다", "필수"),
    ("6", "accuracy", "정확성", "평가", "정확 / 오역있음", "필수"),
    ("7", "fidelity", "충실도", "평가", "충실 / 축소·생략 / 부연·의역", "필수"),
    ("8", "error_history", "오역 이력", "평가", "선행오역수정 / 답습", "필수"),
    ("9", "social_critique", "사회비판 전달", "평가", "살림 / 약화·삭제", "필수"),
    ("10", "interpretive_frame", "해석 프레임", "평가", "연애소설 / 사회비판 / 실존적", "선택"),
]
DIM_NAMES = [d[2] for d in DIMENSIONS]

COLS = ["statement_id", "para_id", "대상번역본", "진술주체", "진술유형",
        "차원", "값", "극성", "근거표현", "요지", "검토", "원문(참조)"]
WIDTHS = [13, 8, 18, 14, 11, 18, 20, 7, 32, 34, 8, 70]
AUTO_COLS = {"statement_id", "para_id", "대상번역본", "원문(참조)"}


def norm_target(t):
    if not t:
        return ""
    t = t.strip()
    if NOT_A_TARGET.match(t):
        return ""
    return TARGET_FIX.get(t, t)


def in_scope(r, all_mode):
    if all_mode:
        return True
    part = r["대분류"] or ""
    sect = r["절"] or ""
    if "일제강점기" in part:
        return True
    if "1945" in part and sect.startswith("2.1"):
        return True
    return False


def load_wonmun(path):
    ws = load_workbook(path)["원문층"]
    hdr = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append(dict(zip(hdr, row)))
    return out


def build_rows(paras, all_mode):
    rows = []
    n = 0
    for p in paras:
        if p["유형"] not in ("서술", "서술(인용포함)"):
            continue
        if not in_scope(p, all_mode):
            continue
        n += 1
        emph = (p["강조구"] or "").split(" | ")[0] if p["강조구"] else ""
        scholars = p["언급연구자"] or ""
        subject = scholars.split(", ")[0] if scholars else "조향"
        rows.append({
            "statement_id": f"CS{n:03d}",
            "para_id": p["para_id"],
            "대상번역본": norm_target(p["대상번역본"]),
            "진술주체": subject,
            "진술유형": "",
            "차원": "",
            "값": "",
            "극성": "",
            "근거표현": emph,
            "요지": "",
            "검토": "",
            "원문(참조)": p["원문"],
        })
    return rows


def style_header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i)
        cell.fill = PatternFill("solid", start_color=PLUM)
        cell.font = Font(bold=True, color="FFFFFF", name=FT, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def write(rows, paras, path, all_mode):
    wb = Workbook()

    # --- 진술 시트 ---
    ws = wb.active
    ws.title = "진술"
    style_header(ws, COLS, WIDTHS)
    for r in rows:
        ws.append([r[c] for c in COLS])
        i = ws.max_row
        for j, c in enumerate(COLS, 1):
            cell = ws.cell(i, j)
            cell.border = BORDER
            cell.font = Font(name=FT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in ("근거표현", "요지", "원문(참조)")))
            cell.fill = PatternFill("solid",
                                    start_color=FILL_AUTO if c in AUTO_COLS else FILL_INPUT)
    ws.freeze_panes = "C2"
    last = ws.max_row

    dv_type = DataValidation(type="list", formula1='"평가,사실,반론"', allow_blank=True)
    dv_dim = DataValidation(type="list", formula1='"' + ",".join(DIM_NAMES) + '"', allow_blank=True)
    dv_pol = DataValidation(type="list", formula1='"+,-,0"', allow_blank=True)
    dv_chk = DataValidation(type="list", formula1='"확인필요,보류,확정"', allow_blank=True)
    for dv in (dv_type, dv_dim, dv_pol, dv_chk):
        ws.add_data_validation(dv)
    dv_type.add(f"E2:E{last + 200}")
    dv_dim.add(f"F2:F{last + 200}")
    dv_pol.add(f"H2:H{last + 200}")
    dv_chk.add(f"K2:K{last + 200}")

    # --- 분류표 시트 ---
    ws2 = wb.create_sheet("분류표")
    style_header(ws2, ["#", "ID", "차원", "구분", "값", "극성"], [5, 20, 20, 8, 46, 8])
    for d in DIMENSIONS:
        ws2.append(list(d))
        i = ws2.max_row
        for j in range(1, 7):
            cell = ws2.cell(i, j)
            cell.border = BORDER
            cell.font = Font(name=FT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.append([])
    guide = [
        ["진술유형", "", "", "", "평가 = 번역의 질에 대한 판단 / 사실 = 저본·경위·생몰년 등 / 반론 = 선행 견해에 대한 이의", ""],
        ["진술주체", "", "", "", "조향 = 비평 저자 본인의 판단 / 그 외 = 인용된 연구자(김미지·이유영·김동소·박진영·조우호·차봉희·김병철)", ""],
        ["극성", "", "", "", "평가 항목(5~10)에만 붙인다. 관찰 항목(1~4)은 비워 둔다.", ""],
        ["한 문단 여러 진술", "", "", "", "행을 복사해 아래에 추가하고 statement_id만 새로 매긴다. para_id는 같아도 된다.", ""],
        ["진술이 없는 문단", "", "", "", "진술유형을 비우고 검토를 '보류'로 둔다. 삭제하지 않는다.", ""],
    ]
    for g in guide:
        ws2.append(g)
        i = ws2.max_row
        ws2.cell(i, 1).font = Font(name=FT, size=10, bold=True)
        for j in range(1, 7):
            ws2.cell(i, j).alignment = Alignment(vertical="top", wrap_text=True)

    # --- 원문참조 시트 ---
    ws3 = wb.create_sheet("원문참조")
    style_header(ws3, ["para_id", "대분류", "절", "대상번역본", "유형", "원문", "강조구"],
                 [8, 14, 22, 18, 13, 86, 34])
    for p in paras:
        if not in_scope(p, all_mode):
            continue
        ws3.append([p["para_id"], p["대분류"], p["절"], norm_target(p["대상번역본"]),
                    p["유형"], p["원문"], p["강조구"]])
        i = ws3.max_row
        for j in range(1, 8):
            cell = ws3.cell(i, j)
            cell.border = BORDER
            cell.font = Font(name=FT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(j in (6, 7)))
    ws3.freeze_panes = "A2"

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("infile", help="UeDeKo_Werther_원문층.xlsx 경로")
    ap.add_argument("-o", "--outdir", default=".")
    ap.add_argument("--all", action="store_true", help="범위 제한 없이 전체 문단 대상")
    args = ap.parse_args()

    paras = load_wonmun(args.infile)
    rows = build_rows(paras, args.all)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    path = outdir / "UeDeKo_Werther_구조층.xlsx"
    write(rows, paras, path, args.all)

    scope = "전체" if args.all else "일제강점기 + 2.1 개관"
    print(f"\n{'='*58}")
    print(f"범위: {scope}")
    print(f"태깅 대상 문단: {len(rows)}개")
    print(f"{'='*58}\n")

    tg = Counter(r["대상번역본"] for r in rows if r["대상번역본"])
    print("[대상번역본별]")
    for k, v in tg.most_common():
        print(f"   {v:>3}  {k}")
    blank = sum(1 for r in rows if not r["대상번역본"])
    print(f"   {blank:>3}  (개관 — 대상 없음)")

    sub = Counter(r["진술주체"] for r in rows)
    print("\n[진술주체 후보]")
    for k, v in sub.most_common():
        print(f"   {v:>3}  {k}")

    emph = sum(1 for r in rows if r["근거표현"])
    print(f"\n[근거표현이 미리 채워진 행] {emph}개")
    print(f"\n저장: {path}\n")


if __name__ == "__main__":
    main()
