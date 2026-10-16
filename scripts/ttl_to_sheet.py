#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ktrans TTL -> 스프레드시트 (사람이 읽는 판)

TTL이 원본이고 시트는 TTL에서 만든다. 그래서 둘이 어긋나지 않는다.
명세 시트(개요~후속과제)는 기존 명세 엑셀을 그대로 쓰고, 사용 횟수와 트리플 수만 TTL로 다시 센다.
그 앞에 '읽는 법'을, 뒤에 실제 데이터 시트를 붙인다.

사용법
    python ttl_to_sheet.py --schema ontology/ktrans-schema.ttl --data data/ktrans-data.ttl \
        --spec docs/ktrans_명세_v3.xlsx -o docs/ktrans_v3.1_데이터.xlsx
"""
import re
import argparse
from collections import Counter, defaultdict
from rdflib import Graph, URIRef, Literal
from rdflib.namespace import RDF, RDFS, SKOS
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

KT = "https://ktrans.aks.ac.kr/id/"
KTR = "https://ktrans.aks.ac.kr/ontology#"
LRM = "http://iflastandards.info/ns/lrm/lrmoo/"
CRM = "http://www.cidoc-crm.org/cidoc-crm/7.1.3/"
PFX = {"ktrans": KTR, "kt": KT, "lrmoo": LRM, "crm": CRM,
       "skos": str(SKOS), "rdfs": str(RDFS)}
SPARQL_PFX = "".join(f"PREFIX {k}: <{v}>\n" for k, v in PFX.items())

PLUM = "6D2E46"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FT = "Arial"


def U(curie):
    p, l = curie.split(":", 1)
    return URIRef(PFX[p] + l)


def short(u):
    s = str(u)
    for k, v in PFX.items():
        if s.startswith(v):
            return f"{k}:{s[len(v):]}"
    return s


def lab(g, node, lang="ko"):
    if node is None:
        return ""
    vals = list(g.objects(node, RDFS.label))
    for v in vals:
        if getattr(v, "language", None) == lang:
            return str(v)
    return str(vals[0]) if vals else short(node)


def note(g, node, lang="ko"):
    return "\n".join(str(o) for o in g.objects(node, U("crm:P3_has_note")) if getattr(o, "language", None) == lang)


def labs(g, node):
    """통제어휘 값: 한국어 · 영어"""
    if node is None:
        return ""
    ko, en = lab(g, node, "ko"), lab(g, node, "en")
    return ko if not en or en == ko else f"{ko} · {en}"


def year(g, node):
    for ts in g.objects(node, U("crm:P4_has_time-span")):
        for y in g.objects(ts, U("crm:P82_at_some_time_within")):
            return str(y)
    return ""


# ---------------------------------------------------------------- 시트 쓰기
def sheet(wb, title, cols, widths, rows, wrap=()):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i)
        cell.fill = PatternFill("solid", start_color=PLUM)
        cell.font = Font(bold=True, color="FFFFFF", name=FT, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for r in rows:
        ws.append(r)
        n = ws.max_row
        for i in range(1, len(cols) + 1):
            cell = ws.cell(n, i)
            cell.font = Font(name=FT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cols[i - 1] in wrap)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    return ws


def build(g, gs, spec, out):
    gall = g + gs
    wb = load_workbook(spec)

    # ---- 1. 명세 시트 숫자 갱신 --------------------------------------
    total = len(gall)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "11,198" in c.value:
                    c.value = c.value.replace("11,198", f"{total:,}")
    changed = []
    for name in ("오브젝트프로퍼티 Object props", "데이터프로퍼티 Data props"):
        ws = wb[name]
        hdr = [str(c.value or "") for c in ws[1]]
        ucol = next(i for i, h in enumerate(hdr) if h.startswith("사용"))
        for row in ws.iter_rows(min_row=2):
            pid = str(row[0].value or "")
            if not re.match(r"^\w+:[\w-]+$", pid):
                continue
            try:
                n = len(set(gall.triples((None, U(pid), None))))
            except KeyError:
                continue
            old = row[ucol].value
            if str(old).isdigit() and int(old) != n:
                changed.append((pid, old, n))
                row[ucol].value = str(n)
    ov = wb["개요 Overview"]
    ov.append([])
    ov.append(["메모 규칙\nNote convention",
               "모든 노드의 설명은 crm:P3_has_note 하나에 담는다. 한국어와 영어를 함께 둔다. 문장 앞에 꼬리표를 붙인다.\n"
               "(설명) 자료에 적힌 사실 · (해석) 누군가의 판단(판단한 사람의 이름을 함께 쓴다) · "
               "(편찬) 데이터를 만든 사람의 판단 · (근거) 문단 번호나 문헌\n"
               "Every node carries one crm:P3_has_note. Each sentence starts with a tag: "
               "(설명) description · (해석) someone's interpretation, named · (편찬) compiler's decision · (근거) evidence."])
    ov.append(["출처\nSource", "UeDeKo — http://www.uedeko.or.kr/wiki2/index.php?title=젊은_베르터의_고뇌_(Die_Leiden_des_jungen_Werthers)\n"
               "서울대학교 독일어문화권연구소·(재)한독문학번역연구소, NRF 2022S1A5C2A02093427"])
    for r in (ov.max_row - 1, ov.max_row):
        for c in ov[r]:
            c.font = Font(name=FT, size=10, bold=(c.column == 1))
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # ---- 2. 번역본 -------------------------------------------------
    stmt_count = Counter(o for s, o in g.subject_objects(U("crm:P140_assigned_attribute_to")))
    rows = []
    for t in sorted(g.subjects(RDF.type, U("ktrans:Translation")), key=lambda x: str(x)):
        ev = next(g.subjects(U("lrmoo:R17_created"), t), None)
        tr = ", ".join(lab(g, p) for p in g.objects(ev, U("crm:P14_carried_out_by"))) if ev else ""
        src = next(g.objects(ev, U("crm:P16_used_specific_object")), None) if ev else None
        cert = ""
        if ev:
            for ty in g.objects(ev, U("crm:P2_has_type")):
                if (ty, SKOS.inScheme, U("kt:scheme-certainty")) in gs:
                    cert = labs(gs, ty)
        nm = g.value(None, U("crm:P1_is_identified_by"), Literal(short(t)))
        nt = note(g, t)
        form = re.search(r"게재 형태는 (.+?)이다", nt)
        form_en = re.search(r"Publication form: (.+?)\.", note(g, t, "en"))
        leiden = re.search(r"Leiden을 '(.+?)'", note(g, nm)) if nm else None
        rows.append([short(t), year(g, t), lab(g, t),
                     ", ".join(str(o) for o in g.objects(t, SKOS.altLabel)),
                     tr, (form.group(1) + " · " + form_en.group(1)) if form else "",
                     ", ".join(labs(gs, o) for o in g.objects(t, U("ktrans:sourceType"))),
                     ", ".join(sorted(labs(gs, o) for o in g.objects(t, U("ktrans:viaType")))),
                     labs(g, src) if src else "", cert,
                     leiden.group(1) if leiden else "",
                     stmt_count.get(t, 0), nt, note(g, t, "en")])
    rows.sort(key=lambda r: (r[1], r[0]))
    C = ["ID", "연도\nYear", "제목\nTitle", "수록서·이형\nBook / variant", "역자\nTranslator", "게재 형태\nForm",
         "분량 축\nCoverage", "경유 축\nRoute", "저본\nSource text", "저본 확실성\nCertainty", "Leiden 역어\nLeiden as",
         "비평 진술 수\nStatements", "메모 (ko)\nNote, Korean", "메모 (en)\nNote, English"]
    sheet(wb, "번역본 Translations", C,
          [9, 7, 24, 20, 12, 16, 14, 18, 26, 14, 9, 9, 70, 70], rows, wrap=tuple(C[2:4] + C[8:9] + C[12:]))

    # ---- 3. 비평 진술 ----------------------------------------------
    rows = []
    for cs in g.subjects(U("crm:P2_has_type"), U("kt:type-criticism")):
        types = list(g.objects(cs, U("crm:P2_has_type")))
        st = next((labs(gs, t) for t in types if (t, SKOS.inScheme, U("kt:scheme-statementType")) in gs), "")
        ft = next((labs(gs, t) for t in types if (t, SKOS.inScheme, U("kt:scheme-factType")) in gs), "")
        tgt = g.value(cs, U("crm:P140_assigned_attribute_to"))
        tgt_s = ""
        if tgt is not None:
            tgt_s = f"{short(tgt)} {lab(g, tgt)}" + (f" ({year(g, tgt)})" if year(g, tgt) else "")
        dim = g.value(cs, U("crm:P177_assigned_property_of_type"))
        val = g.value(cs, U("crm:P141_assigned"))
        sp = g.value(cs, U("crm:P16_used_specific_object"))
        rows.append([short(cs), tgt_s, lab(g, g.value(cs, U("crm:P14_carried_out_by"))), st, ft,
                     labs(gs, dim) if dim else "", labs(gs, val) if val else "",
                     str(g.value(cs, U("ktrans:hasPolarity")) or ""),
                     lab(g, cs, "ko"), lab(g, cs, "en"), str(g.value(sp, U("crm:P1_is_identified_by")) or ""),
                     note(g, cs), note(g, cs, "en")])
    rows.sort(key=lambda r: int(re.sub(r"\D", "", r[10]) or 0))
    C = ["ID", "대상\nTarget", "비평가\nCritic", "진술유형\nType", "사실유형\nFact type", "차원\nDimension",
         "값\nValue", "극성\nPolarity", "요지 (ko)\nSummary, Korean", "요지 (en)\nSummary, English",
         "근거 문단\nParagraph", "메모 (ko)\nNote, Korean", "메모 (en)\nNote, English"]
    sheet(wb, "비평 진술 Statements", C,
          [12, 28, 8, 14, 16, 20, 18, 8, 36, 40, 9, 50, 50], rows, wrap=(C[1], C[8], C[9], C[11], C[12]))

    # ---- 4. 인물 ---------------------------------------------------
    ROLE_EN = {"저자": "author", "역자": "translator", "비평가·연구자": "critic or scholar"}
    role = defaultdict(set)
    for s, p in g.subject_objects(U("crm:P14_carried_out_by")):
        if (s, RDF.type, U("lrmoo:F27_Work_Creation")) in g:
            role[p].add("저자")
        elif (s, RDF.type, U("lrmoo:F28_Expression_Creation")) in g:
            role[p].add("역자")
        else:
            role[p].add("비평가·연구자")
    rows = []
    for p in sorted(g.subjects(RDF.type, U("crm:E21_Person")), key=lambda x: lab(g, x)):
        rows.append([short(p), lab(g, p), ", ".join(f"{r} · {ROLE_EN[r]}" for r in sorted(role[p])),
                     ", ".join(sorted(str(o) for o in g.objects(p, SKOS.altLabel))),
                     "\n".join(str(o) for o in g.objects(p, U("crm:P1_is_identified_by"))),
                     note(g, p), note(g, p, "en")])
    C = ["ID", "정규형\nName", "역할\nRole", "표기 이형\nVariant forms (skos:altLabel)", "외부 전거\nAuthority links",
         "메모 (ko)\nNote, Korean", "메모 (en)\nNote, English"]
    sheet(wb, "인물 Persons", C, [16, 14, 18, 40, 40, 50, 50], rows, wrap=tuple(C[3:]))

    # ---- 5. 원문 문단 ----------------------------------------------
    used = Counter(o for o in g.objects(None, U("crm:P16_used_specific_object")))
    rows = []
    for sp in g.subjects(U("crm:P2_has_type"), U("kt:type-source-paragraph")):
        pid = str(g.value(sp, U("crm:P1_is_identified_by")))
        rows.append([pid, short(sp), used.get(sp, 0), str(g.value(sp, U("crm:P190_has_symbolic_content")))])
    rows.sort(key=lambda r: int(r[0].split("_")[1]))
    C = ["para_id", "ID", "근거로 쓴 진술 수\nStatements citing it", "본문 (UeDeKo 번역비평 원문, 한국어)\nText of the UeDeKo critical essay (Korean original)"]
    sheet(wb, "원문 문단 Paragraphs", C, [9, 8, 12, 120], rows, wrap=(C[3],))

    # ---- 6. 읽는 법 -------------------------------------------------
    rd = wb.create_sheet("읽는 법 README", 0)
    n_t = sum(1 for _ in g.subjects(RDF.type, U('ktrans:Translation')))
    n_s = len(set(g.subjects(U('crm:P2_has_type'), U('kt:type-criticism'))))
    lines = [
        ["ktrans — 괴테 『젊은 베르터의 고뇌』 한국어 번역·비평 링크드 데이터 (v3.1)"],
        ["ktrans — Korean translations of Goethe's Werther and their criticism as Linked Data (v3.1)"],
        [],
        ["이 파일은\nThis file", "Turtle 파일(ktrans-schema.ttl, ktrans-data.ttl)을 사람이 읽을 수 있게 표로 펼친 것이다. 원본은 TTL이고, 이 표는 스크립트(ttl_to_sheet.py)로 TTL에서 만든다.\n"
                       "The Turtle files laid out as tables for people. The TTL is the source; this workbook is generated from it by ttl_to_sheet.py."],
        ["원 자료\nSource", "UeDeKo 『젊은 베르터의 고뇌』 페이지의 번역서지 표(318행)와 번역비평 줄글(조향). 전공자들이 모으고 텍스트화한, 사람이 읽는 자료를 기계와 AI가 읽는 RDF로 옮겼다.\n"
                     "The bibliography table (318 rows) and the critical essay (Jo Hyang) on the UeDeKo Werther page: material gathered by specialists for people, converted into RDF that machines and AI can read."],
        ["작성\nCompiled by", "정송이 (Jung Song-yi), 한국학중앙연구원 디지털인문학연구소 · Digital Humanities Research Institute, Academy of Korean Studies"],
        [],
        ["시트 안내\nSheets", ""],
        ["  개요 ~ 후속과제", "온톨로지 명세. 쓰인 클래스·프로퍼티·통제어휘와 설계 결정.\nOverview to Open items: the ontology specification — classes, properties, vocabularies and design decisions."],
        ["  번역본 Translations", f"작품 단위 번역 {n_t}건. UeDeKo 레코드 318건을 연재·분책·중복등재 통합해 만든 단위.\n{n_t} translations at work level, merged from 318 UeDeKo records (serials, split volumes, duplicates)."],
        ["  비평 진술 Statements", f"번역비평을 진술 단위로 나눈 {n_s}건. 누가·무엇을·어떤 기준으로·어떤 극성으로·어느 문단에 근거해.\n{n_s} critical statements: who, about which translation, on which dimension, with which polarity, from which paragraph."],
        ["  인물 Persons", "괴테, 역자, 비평가·연구자. 역할은 클래스가 아니라 참여한 사건에서 나온다.\nGoethe, translators, critics and scholars. A role comes from the event a person takes part in, not from a class."],
        ["  원문 문단 Paragraphs", "UeDeKo 번역비평 원문을 문단 단위로 나눈 것. 모든 진술의 근거.\nThe UeDeKo critical essay split into paragraphs: the evidence for every statement. Korean original only."],
        [],
        ["메모 규칙\nNote convention", "'메모' 칸은 TTL의 crm:P3_has_note 그대로다(한국어 @ko, 영어 @en). 문장 앞 꼬리표로 성격을 구분한다.\n"
                                "The note columns are crm:P3_has_note as stored, in Korean (@ko) and English (@en). Each sentence starts with a tag."],
        ["  (설명) description", "자료에 적힌 사실. UeDeKo 서지·비고, 게재 형태, 역어.\nWhat the source says: UeDeKo's bibliography and remarks, publication form, title word."],
        ["  (해석) interpretation", "누군가의 판단. 판단한 사람의 이름을 문장 안에 쓴다. 예: 김동소는 … 추정한다.\nSomeone's judgement, with the person named in the sentence, e.g. 'Kim Dong-so infers …'."],
        ["  (편찬) compilation", "데이터를 만든 사람(정송이)의 판단. 레코드 통합, 이름 정규화, 비교 기준.\nThe compiler's decision: merging records, normalising names, stating the basis of comparison."],
        ["  (근거) evidence", "문단 번호(para_n)나 문헌.\nA paragraph number (para_n) or a reference."],
        [],
        ["읽는 예\nExample", "번역본 T005(백화 역, 1923) → 메모에서 게재 형태(신문 연재)와 통합 근거를 읽고, '비평 진술' 시트에서 대상이 kt:T005인 행을 필터하면 그 번역에 대한 진술 16건과 근거 문단이 나온다.\n"
                       "Translation T005 (Baek-hwa, 1923): its note gives the publication form (newspaper serial) and the reason for merging; filter the Statements sheet for target kt:T005 to see its 16 statements and their paragraphs."],
        ["중역과 충실도\nRelay and fidelity", "중역 번역의 충실도 진술 12건에는 비교 기준(독일어 원문)을 밝히는 (편찬) 문장이 있다. 생략이 중간 저본에서 왔는지는 데이터로 알 수 없다.\n"
                       "The 12 fidelity statements on relay translations carry a (compilation) sentence naming the basis of comparison (the German original); the data cannot tell whether an omission comes from the intermediate source."],
        ["영어 문장\nEnglish text", "영어 메모와 진술 요지는 편찬자가 한국어 원문에서 옮긴 것이다. 근거 인용과 원문 문단은 한국어 원문 그대로 둔다.\n"
                       "English notes and statement summaries were translated from the Korean by the compiler. Quotations and paragraphs stay in the Korean original."],
        [],
        ["출처\nCredit", "UeDeKo http://www.uedeko.or.kr/wiki2/index.php?title=젊은_베르터의_고뇌_(Die_Leiden_des_jungen_Werthers) — 서울대학교 독일어문화권연구소·(재)한독문학번역연구소 · Institute of German Studies and Korean-German Literary Translation Institute, Seoul National University (NRF 2022S1A5C2A02093427)"],
        ["트리플 수\nTriples", f"{total:,} (schema {len(gs):,} + data {len(g):,})"],
    ]
    for r in lines:
        rd.append(r)
    rd.column_dimensions["A"].width = 24
    rd.column_dimensions["B"].width = 110
    for row in rd.iter_rows():
        for c in row:
            c.font = Font(name=FT, size=10, bold=(c.column == 1 and not str(c.value or "").startswith("  ")))
            c.alignment = Alignment(vertical="top", wrap_text=True)
    rd["A1"].font = Font(name=FT, size=13, bold=True, color=PLUM)
    rd["A2"].font = Font(name=FT, size=10, italic=True, color="666666")
    wb.active = 0
    wb.save(out)
    return changed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--schema", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--spec", required=True)
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    gs = Graph(); gs.parse(a.schema)
    g = Graph(); g.parse(a.data)
    changed = build(g, gs, a.spec, a.out)
    for pid, old, n in changed:
        print(f"  사용 횟수 갱신: {pid} {old} -> {n}")
    print("저장:", a.out)


if __name__ == "__main__":
    main()
